"""
Módulo de exportación: Excel y reportes de texto formateados.
"""

import os
import subprocess
import platform
from datetime import datetime
from typing import List, Dict, Optional


def exportar_a_excel(columnas: List[str], datos: List[Dict], nombre_tabla: str) -> Optional[str]:
    """
    Exporta datos a un archivo Excel (.xlsx).
    Retorna la ruta del archivo generado o None si falla.
    """
    try:
        import openpyxl
        from openpyxl.styles import (
            PatternFill, Font, Alignment, Border, Side
        )

        # Crear workbook
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = nombre_tabla[:31]  # Excel limita a 31 chars

        # Estilos
        header_fill = PatternFill(start_color="1C4E80", end_color="1C4E80", fill_type="solid")
        header_font = Font(color="FFFFFF", bold=True, size=11)
        alt_fill = PatternFill(start_color="F0F4F8", end_color="F0F4F8", fill_type="solid")
        border_side = Side(style="thin", color="CCCCCC")
        border = Border(
            left=border_side, right=border_side,
            top=border_side, bottom=border_side
        )
        center_align = Alignment(horizontal="center", vertical="center")

        # Título de la hoja
        ws.merge_cells("A1:{}1".format(chr(64 + len(columnas))))
        title_cell = ws["A1"]
        title_cell.value = f"GESTIÓN SOLAR DEL CARIBE S.A.S. — {nombre_tabla}"
        title_cell.font = Font(bold=True, size=13, color="1C4E80")
        title_cell.alignment = center_align
        title_cell.fill = PatternFill(start_color="E8F4FD", end_color="E8F4FD", fill_type="solid")
        ws.row_dimensions[1].height = 28

        # Fecha de exportación
        ws.merge_cells("A2:{}2".format(chr(64 + len(columnas))))
        date_cell = ws["A2"]
        date_cell.value = f"Exportado el: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}"
        date_cell.font = Font(italic=True, size=10, color="666666")
        date_cell.alignment = center_align

        # Encabezados (fila 4)
        for col_idx, col_name in enumerate(columnas, start=1):
            cell = ws.cell(row=4, column=col_idx, value=col_name.replace("_", " ").title())
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = center_align
            cell.border = border
        ws.row_dimensions[4].height = 22

        # Datos
        for row_idx, row_data in enumerate(datos, start=5):
            fill = alt_fill if row_idx % 2 == 0 else None
            for col_idx, col_name in enumerate(columnas, start=1):
                value = row_data.get(col_name, "")
                # Convertir booleanos
                if isinstance(value, int) and col_name in ("limpieza_paneles", "activo"):
                    value = "Sí" if value else "No"
                cell = ws.cell(row=row_idx, column=col_idx, value=value)
                cell.border = border
                cell.alignment = Alignment(vertical="center", wrap_text=True)
                if fill:
                    cell.fill = fill

        # Ajustar anchos de columna automáticamente
        for col_idx, col_name in enumerate(columnas, start=1):
            col_letter = chr(64 + col_idx) if col_idx <= 26 else "A" + chr(64 + col_idx - 26)
            max_length = max(
                len(str(col_name)),
                max((len(str(row.get(col_name, "") or "")) for row in datos), default=0)
            )
            ws.column_dimensions[col_letter].width = min(max_length + 4, 40)

        # Guardar archivo
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        filename = f"Solar_Caribe_{nombre_tabla}_{timestamp}.xlsx"
        filepath = os.path.join(os.path.expanduser("~"), "Documents", filename)

        # Crear carpeta Documents si no existe
        os.makedirs(os.path.dirname(filepath), exist_ok=True)

        wb.save(filepath)
        return filepath

    except ImportError:
        return None
    except Exception as e:
        print(f"Error exportando a Excel: {e}")
        return None


def generar_informe_mantenimiento(mantenimiento: Dict, proyecto: Dict, cliente: Dict) -> Optional[str]:
    """
    Genera un informe de mantenimiento formateado en texto.
    Retorna la ruta del archivo generado.
    """
    try:
        timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
        nombre_cliente = cliente.get("nombre_titular", "Cliente").replace(" ", "_")
        filename = f"Informe_Mant_{nombre_cliente}_{timestamp}.txt"
        filepath = os.path.join(os.path.expanduser("~"), "Documents", filename)
        os.makedirs(os.path.dirname(filepath), exist_ok=True)

        # Calidad de estado
        estado_labels = {1: "Crítico", 2: "Deficiente", 3: "Regular", 4: "Bueno", 5: "Excelente"}

        lineas = [
            "=" * 65,
            "         GESTIÓN SOLAR DEL CARIBE S.A.S.",
            "         NIT: XXX.XXX.XXX-X",
            "         Tel: +57 (605) XXX-XXXX",
            "         solar.caribe@ejemplo.com",
            "=" * 65,
            "",
            "              INFORME DE MANTENIMIENTO TÉCNICO",
            "",
            "-" * 65,
            "  INFORMACIÓN DEL CLIENTE",
            "-" * 65,
            f"  Titular:         {cliente.get('nombre_titular', 'N/A')}",
            f"  Cédula/NIT:      {cliente.get('cedula_nit', 'N/A')}",
            f"  Dirección:       {cliente.get('direccion_completa', 'N/A')}",
            f"  Barrio/Sector:   {cliente.get('barrio_sector', 'N/A')}",
            f"  Estrato:         {cliente.get('estrato', 'N/A')}",
            f"  Teléfono:        {cliente.get('telefono', 'N/A')}",
            f"  Email:           {cliente.get('email', 'N/A')}",
            "",
            "-" * 65,
            "  DATOS DEL SISTEMA SOLAR INSTALADO",
            "-" * 65,
            f"  Capacidad:       {proyecto.get('kwp_instalados', 'N/A')} kWp",
            f"  Fecha Inst.:     {proyecto.get('fecha_instalacion', 'N/A')}",
            f"  Marca Paneles:   {proyecto.get('marca_paneles', 'N/A')}",
            f"  Seriales:        {proyecto.get('seriales_paneles', 'N/A')}",
            f"  Inversor:        {proyecto.get('marca_inversor', 'N/A')} {proyecto.get('capacidad_inversor_kw', '')} kW",
            f"  Serial Inversor: {proyecto.get('serial_inversor', 'N/A')}",
            f"  Estado RETIE:    {proyecto.get('estatus_retie', 'N/A')}",
            "",
            "-" * 65,
            "  REGISTRO DE VISITA DE MANTENIMIENTO",
            "-" * 65,
            f"  Fecha Visita:    {mantenimiento.get('fecha_visita', 'N/A')}",
            f"  Técnico:         {mantenimiento.get('tecnico_encargado', 'N/A')}",
            f"  Fotos Evidencia: {mantenimiento.get('fotos_evidencia_path', 'No registrado')}",
            "",
            "-" * 65,
            "  RESULTADOS DE INSPECCIÓN",
            "-" * 65,
            f"  Estado Estructura:  {mantenimiento.get('estado_estructura', 'N/A')}/5 "
            f"— {estado_labels.get(mantenimiento.get('estado_estructura', 0), 'N/A')}",
            f"  Estado Cableado:    {mantenimiento.get('estado_cableado', 'N/A')}/5 "
            f"— {estado_labels.get(mantenimiento.get('estado_cableado', 0), 'N/A')}",
            f"  Limpieza Paneles:   {'✓ Realizada' if mantenimiento.get('limpieza_paneles') else '✗ No realizada'}",
            f"  Continuidad Tierra: {mantenimiento.get('continuidad_tierra_ohm', 'N/A')} Ω",
            "",
            "-" * 65,
            "  OBSERVACIONES CRÍTICAS",
            "-" * 65,
        ]

        observaciones = mantenimiento.get("observaciones_criticas", "Sin observaciones críticas.")
        # Partir observaciones en líneas de máx 60 chars
        if observaciones:
            palabras = observaciones.split()
            linea_actual = "  "
            for palabra in palabras:
                if len(linea_actual) + len(palabra) + 1 > 62:
                    lineas.append(linea_actual)
                    linea_actual = "  " + palabra + " "
                else:
                    linea_actual += palabra + " "
            if linea_actual.strip():
                lineas.append(linea_actual)
        else:
            lineas.append("  Sin observaciones críticas.")

        lineas += [
            "",
            "=" * 65,
            f"  Generado: {datetime.now().strftime('%d/%m/%Y %H:%M:%S')}",
            f"  Por: Sistema GSC v1.0",
            "=" * 65,
            "",
            "  Firma Técnico: _______________________",
            "",
            "  Firma Cliente: _______________________",
            "",
        ]

        with open(filepath, "w", encoding="utf-8") as f:
            f.write("\n".join(lineas))

        return filepath

    except Exception as e:
        print(f"Error generando informe: {e}")
        return None


def abrir_carpeta_windows(ruta: str) -> bool:
    """Abre una carpeta en el explorador de Windows."""
    if not ruta or not os.path.exists(ruta):
        return False
    try:
        if platform.system() == "Windows":
            os.startfile(ruta)
        elif platform.system() == "Darwin":
            subprocess.Popen(["open", ruta])
        else:
            subprocess.Popen(["xdg-open", ruta])
        return True
    except Exception:
        return False


def abrir_archivo(ruta: str) -> bool:
    """Abre un archivo con la aplicación predeterminada."""
    if not ruta or not os.path.exists(ruta):
        return False
    try:
        if platform.system() == "Windows":
            os.startfile(ruta)
        elif platform.system() == "Darwin":
            subprocess.Popen(["open", ruta])
        else:
            subprocess.Popen(["xdg-open", ruta])
        return True
    except Exception:
        return False
